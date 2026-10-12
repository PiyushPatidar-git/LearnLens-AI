import os, base64, re, time
from .db import conn
from .text import keyphrase
from . import llm
EXT = {".pdf": "pdf", ".pptx": "slide", ".docx": "docx", ".png": "image", ".jpg": "image", ".jpeg": "image",
       ".mp4": "video", ".mkv": "video", ".mov": "video", ".webm": "video", ".mp3": "video", ".wav": "video"}
def detect(name): return EXT.get(os.path.splitext(name)[1].lower())
def _ocr(img_bytes):
    try:
        import pytesseract, io
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(img_bytes)))
    except Exception: return ""
def _caption(img_bytes):
    """VLM diagram understanding; falls back to OCR text if no VLM available."""
    r = llm.complete("You describe educational diagrams/charts/tables precisely for a study knowledge base.",
                     "Describe this figure: what it shows, labels, relationships, and the concept it illustrates.",
                     400, base64.b64encode(img_bytes).decode())
    return ("[Figure] " + r) if r else (("[Figure OCR] " + _ocr(img_bytes)) if _ocr(img_bytes).strip() else "")
def _split(text, n=900):
    paras, cur, out = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()], "", []
    for p in paras:
        if len(cur) + len(p) > n and cur: out.append(cur); cur = ""
        cur += p + "\n"
    if cur.strip(): out.append(cur)
    return out
def parse_pdf(path):
    import fitz
    doc, units = fitz.open(path), []
    for i, pg in enumerate(doc, 1):
        t = pg.get_text()
        if len(t.strip()) < 40:  # scanned page -> OCR
            t = _ocr(pg.get_pixmap(dpi=150).tobytes("png"))
        for ch in _split(t): units.append(dict(text=ch, kind="page", page=i))
        for img in pg.get_images(full=True)[:3]:
            try:
                cap = _caption(doc.extract_image(img[0])["image"])
                if cap: units.append(dict(text=cap, kind="page", page=i))
            except Exception: pass
    return units
def parse_pptx(path):
    from pptx import Presentation
    units = []
    for i, s in enumerate(Presentation(path).slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame: parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.append("\n".join(" | ".join(c.text for c in r.cells) for r in sh.table.rows))
            if sh.shape_type == 13:
                cap = _caption(sh.image.blob)
                if cap: parts.append(cap)
        if s.has_notes_slide: parts.append(s.notes_slide.notes_text_frame.text)
        t = "\n".join(p for p in parts if p.strip())
        if t.strip(): units.append(dict(text=t, kind="slide", slide=i))
    return units
def parse_docx(path):
    import docx
    return [dict(text=c, kind="page", page=1) for c in _split("\n\n".join(p.text for p in docx.Document(path).paragraphs))]
def parse_image(path):
    b = open(path, "rb").read(); t = _caption(b)
    return [dict(text=t, kind="page", page=1)] if t.strip() else []
def parse_video(path, window=60):
    from faster_whisper import WhisperModel
    segs, _ = WhisperModel("base", compute_type="int8").transcribe(path)
    units, cur, start, end = [], "", 0, 0
    for s in segs:
        if not cur: start = s.start
        cur += " " + s.text.strip(); end = s.end
        if end - start >= window: units.append(dict(text=cur.strip(), kind="video", start=start, end=end)); cur = ""
    if cur.strip(): units.append(dict(text=cur.strip(), kind="video", start=start, end=end))
    return units
PARSERS = {"pdf": parse_pdf, "slide": parse_pptx, "docx": parse_docx, "image": parse_image, "video": parse_video}
def run(source_id):
    c = conn(); s = c.execute("select * from sources where id=?", (source_id,)).fetchone()
    try:
        c.execute("update sources set status='extracting' where id=?", (source_id,)); c.commit()
        units = PARSERS[s["type"]](s["path"])
        c.execute("update sources set status='identifying topics' where id=?", (source_id,)); c.commit()
        for u in units:
            topic = keyphrase(u["text"], 2)
            llm_t = llm.complete_json("Extract study metadata.", f'Text:\n{u["text"][:1200]}\nJSON keys: topic, subtopic, concept (short noun phrases).', 200)
            sub = (llm_t or {}).get("subtopic") or keyphrase(u["text"], 3)
            con = (llm_t or {}).get("concept") or keyphrase(u["text"], 1)
            topic = (llm_t or {}).get("topic") or topic
            c.execute("insert into source_chunks(source_id,course_id,text,locator_kind,page,slide,start_s,end_s,topic,subtopic,concept) values(?,?,?,?,?,?,?,?,?,?,?)",
                      (source_id, s["course_id"], u["text"], u["kind"], u.get("page"), u.get("slide"), u.get("start"), u.get("end"), topic, sub, con))
        c.execute("update sources set status='complete' where id=?", (source_id,))
        _prereqs(c, s["course_id"])
    except ModuleNotFoundError as e:
        c.execute("update sources set status='failed', error=? where id=?", (f"missing optional dependency: {e.name}", source_id))
    except Exception as e:
        c.execute("update sources set status='failed', error=? where id=?", (str(e)[:300], source_id))
    c.commit(); c.close()
def _prereqs(c, course_id):
    """Heuristic prerequisite graph: topics ordered by first appearance in the course; each requires the previous one
    (+ LLM-refined edges when available)."""
    rows = c.execute("select topic, min(id) m from source_chunks where course_id=? group by topic order by m", (course_id,)).fetchall()
    names = [r["topic"] for r in rows]
    c.execute("delete from prerequisites where course_id=?", (course_id,))
    edges = None
    if llm.available() and len(names) > 1:
        edges = llm.complete_json("You build prerequisite graphs.", f"Topics: {names}\nReturn a JSON list of [topic, requires] pairs using only these names.", 800)
    if not edges: edges = [[names[i], names[i-1]] for i in range(1, len(names))]
    for t, r in edges:
        if t in names and r in names: c.execute("insert into prerequisites(course_id,topic,requires) values(?,?,?)", (course_id, t, r))
